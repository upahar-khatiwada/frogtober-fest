"""Company document parsing: PDF / DOCX / TXT bytes -> clean Unicode text.

Devanagari-safe: text is decoded as UTF-8 and NFC-normalised, and only
control characters are removed, never non-ASCII. No OCR (out of scope): a
PDF with no extractable text is returned with `needs_ocr=True` rather than
as an empty success.
"""

import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import PurePath

import docx
import pymupdf

# A page with fewer non-whitespace characters than this counts as "no text"
# (a scan, or a page holding only a stamp/signature image).
MIN_PAGE_CHARS = 20

# Legacy Nepali fonts map Devanagari glyphs onto ASCII codepoints, so the
# extracted "text" is Latin garbage (e.g. Preeti renders नेपाल from "g]kfn").
LEGACY_NEPALI_FONTS = ("preeti", "kantipur", "sagarmatha", "himali", "fontasy", "pcs nepali")

# Devanagari glued to Latin-1/Latin-Extended letters inside one word is the
# signature of conjunct glyphs with a broken ToUnicode map (e.g. प्र -> "Ś").
_GARBLED_DEVANAGARI = re.compile(r"[ऀ-ॿ][À-ɏ]|[À-ɏ][ऀ-ॿ]")

PDF_TYPES = {".pdf": "pdf"}
DOCX_TYPES = {".docx": "docx"}
TEXT_TYPES = {".txt": "text", ".md": "text"}
SUPPORTED = {**PDF_TYPES, **DOCX_TYPES, **TEXT_TYPES}


class ParseError(ValueError):
    """The upload cannot be parsed (unsupported type, corrupt file, not UTF-8)."""


@dataclass
class ParsedDocument:
    text: str
    kind: str  # pdf | docx | text
    page_count: int | None = None
    needs_ocr: bool = False
    warnings: list[str] = field(default_factory=list)


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​﻿]")
_SPACES = re.compile(r"[ \t ]+")
_BLANKS = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """NFC-normalise, drop control chars, tidy whitespace. Keeps all scripts.

    ZWJ/ZWNJ (U+200D/U+200C) are kept: they change Devanagari conjunct rendering.
    """
    text = unicodedata.normalize("NFC", text).replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


def detect_kind(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ParseError(f"unsupported file type {suffix or '(none)'!r}; upload PDF, DOCX or TXT")
    return SUPPORTED[suffix]


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # pymupdf raises its own FileDataError family
        raise ParseError(f"not a readable PDF: {exc}") from exc
    with pdf:
        if pdf.needs_pass:
            raise ParseError("PDF is password-protected")
        pages = [clean_text(page.get_text("text")) for page in pdf]
        fonts = {font[3].lower() for page in pdf for font in page.get_fonts()}

    empty = [i + 1 for i, page in enumerate(pages) if len(re.sub(r"\s", "", page)) < MIN_PAGE_CHARS]
    doc = ParsedDocument(text="\n\n".join(p for p in pages if p), kind="pdf", page_count=len(pages))
    if len(empty) == len(pages):
        doc.needs_ocr = True
        doc.text = ""
        doc.warnings.append("no extractable text: looks like a scanned PDF (OCR is not supported)")
    elif empty:
        doc.warnings.append(f"pages without extractable text (scanned?): {', '.join(map(str, empty))}")
    if any(legacy in font for font in fonts for legacy in LEGACY_NEPALI_FONTS):
        doc.warnings.append("uses a legacy Nepali font (e.g. Preeti); extracted Devanagari may be garbled")
    if looks_garbled_devanagari(doc.text):
        doc.warnings.append("some Devanagari conjuncts did not extract cleanly (PDF lacks a Unicode mapping)")
    return doc


def looks_garbled_devanagari(text: str) -> bool:
    return bool(_GARBLED_DEVANAGARI.search(text))


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises zipfile/KeyError/ValueError variants
        raise ParseError(f"not a readable DOCX: {exc}") from exc
    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # Merged cells repeat; keep each distinct value once, in order.
            blocks.append(" | ".join(dict.fromkeys(c for c in cells if c)))
    return ParsedDocument(text=clean_text("\n".join(blocks)), kind="docx")


def parse_text(data: bytes) -> ParsedDocument:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raw = data.decode("utf-16")
    else:
        try:
            raw = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ParseError(f"text file is not UTF-8 (byte {exc.start}); re-save it as UTF-8") from exc
    return ParsedDocument(text=clean_text(raw), kind="text")


def parse_document(data: bytes, filename: str) -> ParsedDocument:
    kind = detect_kind(filename)
    doc = {"pdf": parse_pdf, "docx": parse_docx, "text": parse_text}[kind](data)
    if not doc.text and not doc.needs_ocr:
        doc.warnings.append("document is empty")
    return doc
