import io

import docx
import pymupdf
import pytest

from ingest.parse_docs import ParseError, clean_text, looks_garbled_devanagari, parse_document

NEPALI = "कर चुक्ता प्रमाणपत्र: आ.व. २०८१/८२, पान नं. ९९९०००१११"


def make_pdf(pages: list[str | None]) -> bytes:
    """Each entry is a page: text, or None for an image-only (scanned-looking) page."""
    pdf = pymupdf.open()
    for content in pages:
        page = pdf.new_page()
        if content is None:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
            pix.clear_with(180)
            page.insert_image(pymupdf.Rect(50, 50, 500, 700), pixmap=pix)
        else:
            page.insert_text((72, 72), content)
    data = pdf.tobytes()
    pdf.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_txt_keeps_devanagari():
    doc = parse_document(NEPALI.encode("utf-8"), "tax.txt")
    assert doc.text == NEPALI
    assert doc.kind == "text" and not doc.needs_ocr


def test_txt_with_bom_and_utf16():
    assert parse_document(b"\xef\xbb\xbf" + NEPALI.encode(), "a.txt").text == NEPALI
    assert parse_document(NEPALI.encode("utf-16"), "a.txt").text == NEPALI


def test_txt_not_utf8_is_an_error_not_mojibake():
    with pytest.raises(ParseError, match="not UTF-8"):
        parse_document("Café".encode("latin-1"), "a.txt")


def test_docx_paragraphs_and_tables():
    data = make_docx(["Experience certificate", NEPALI], table=[["Project", "Value"], ["Road gravelling", "रु. १,८५,००,०००"]])
    doc = parse_document(data, "exp.DOCX")
    assert "Experience certificate" in doc.text
    assert NEPALI in doc.text
    assert "Road gravelling | रु. १,८५,००,०००" in doc.text


def test_pdf_text_is_extracted():
    doc = parse_document(make_pdf(["Tax Clearance Certificate FY 2081/82", "Valid until end of FY 2082/83"]), "tc.pdf")
    assert doc.page_count == 2
    assert "Tax Clearance Certificate FY 2081/82" in doc.text
    assert "Valid until" in doc.text
    assert not doc.needs_ocr and doc.warnings == []


def test_scanned_pdf_is_flagged_not_silently_empty():
    doc = parse_document(make_pdf([None, None]), "scan.pdf")
    assert doc.needs_ocr is True
    assert doc.text == ""
    assert doc.page_count == 2
    assert any("scanned" in w for w in doc.warnings)


def test_partly_scanned_pdf_warns_with_page_numbers():
    doc = parse_document(make_pdf(["Company profile of the demo firm, page one.", None]), "mixed.pdf")
    assert not doc.needs_ocr
    assert doc.warnings == ["pages without extractable text (scanned?): 2"]


def test_corrupt_and_unsupported_files():
    with pytest.raises(ParseError, match="PDF"):
        parse_document(b"not a pdf", "x.pdf")
    with pytest.raises(ParseError, match="DOCX"):
        parse_document(b"not a zip", "x.docx")
    with pytest.raises(ParseError, match="unsupported"):
        parse_document(b"...", "x.jpg")


def test_clean_text_keeps_scripts_and_joiners():
    raw = "क्‍ष\x00 \t नेपाल\r\n\r\n\r\n\r\nline"
    assert clean_text(raw) == "क्‍ष नेपाल\n\nline"
    # NFC: U+0958 is a composition exclusion, so its canonical form is base + nukta (not dropped).
    assert clean_text("क़") == "क़"


def test_garbled_devanagari_heuristic():
    # What PyMuPDF extracts from a PDF whose conjunct glyphs have no ToUnicode entry.
    assert looks_garbled_devanagari("कर चुüता Śमाणपŕ")
    assert not looks_garbled_devanagari("कर चुक्ता प्रमाणपत्र — Café, São Paulo")
