"""Bid package export: checklist, package contents, Devanagari PDF (shaping + text extraction), DOCX, the ZIP route.

Pure tests need nothing (fonts are vendored in export/fonts/); the package-route test needs Postgres (pgserver).
"""

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime, timedelta

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from sqlalchemy import select

from db.models import Document as DocumentRow
from db.models import ProposalRow
from export.content import BidPackage, Block, build_checklist, build_package, markdown_blocks, para
from export.docx import _RPR_ORDER, render_docx
from export.package import build
from export.pdf import render_pdf
from nepal.rules import DocumentMeta, evaluate_rules
from schemas import (ChecklistStatus, ComplianceView, Proposal, ProposalSection, RequiredDoc, ReviewReport, ReviewScores)
from tests.test_compliance import ROOT, golden_record
from tests.test_matching import COMPANY, NOW, record, settings_
from tests.test_proposal import db_client
from tests.test_review import ReviewFake, drafted_001, versions
from tests.test_tenders_list_api import api  # noqa: F401  (fixture)

P, E, M = ChecklistStatus.PRESENT, ChecklistStatus.EXPIRING, ChecklistStatus.MISSING
DEVANAGARI_WORD = re.compile(r"[ऀ-ॿ‌‍]+")
PUA = re.compile(r"[-]")


def norm(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).split())


def pdf_text(data: bytes, *, body_only: bool = False) -> str:
    """Extracted text; `body_only` leaves out the footer (the bottom 15 mm of each page)."""
    pages = pymupdf.open(stream=data, filetype="pdf")
    return "".join(page.get_text(clip=pymupdf.Rect(0, 0, page.rect.width, page.rect.height - 42.5) if body_only else None)
                   for page in pages)


# The demo KB as Compliance sees it (FICTIONAL, data/demo_company/manifest.json), with the licence copy expiring
# inside 001's bid validity period (deadline 2026-11-21 + 90 days = 2027-02-19) to exercise "expiring".
DOCS = [
    DocumentMeta(doc_type="company_registration", chunk_id=1, document_id=1, filename="01_company_registration.txt",
                 expiry=date(2027, 7, 16)),
    DocumentMeta(doc_type="pan_vat", chunk_id=3, document_id=2, filename="02_pan_vat_certificate.docx"),
    DocumentMeta(doc_type="tax_clearance", chunk_id=5, document_id=3, filename="03_tax_clearance_2081-82.pdf",
                 expiry=date(2026, 7, 16), fiscal_year_bs="2081/82"),
    DocumentMeta(doc_type="experience_certificate", chunk_id=7, document_id=4, filename="04_experience_certificate.docx"),
    DocumentMeta(doc_type="business_licence", chunk_id=9, document_id=6, filename="06_construction_licence.txt",
                 expiry=date(2027, 1, 10), fiscal_year_bs="2083/84"),
]
R001 = golden_record("001")  # mixed notice -> Nepali proposal; asks for FY 2082/83 tax clearance, bid security, JV optional
ITEMS = evaluate_rules(R001, COMPANY, DOCS, now=NOW)

SECTIONS_NE = [
    ProposalSection(key="cover_letter", title="आवरण पत्र", grounded=True, cited_chunk_ids=[1],
                    body="**अन्नपूर्ण डेमो निर्माण सेवा प्रा. लि.** ले भानु नगरपालिकाको सडक स्तरोन्नतिको बोलपत्र पेस गर्दछ।"),
    ProposalSection(key="relevant_experience", title="सम्बन्धित अनुभव", grounded=True, cited_chunk_ids=[7, 9],
                    body="- व्यास नगरपालिकाको सडक ग्राभेल कार्य, रु. १,८५,४०,०००\n- क्षेत्रीय सडक निर्माण र कल्भर्ट\n\n"
                         "यो कार्य सम्झौता अनुसार सम्पन्न भयो; श्री प्रमुख प्रशासकीय अधिकृतद्वारा प्रमाणित।"),
    ProposalSection(key="financial_capacity", title="आर्थिक क्षमता", grounded=False,
                    flags=[{"kind": "ungrounded", "detail": "x"}], body="औसत वार्षिक कारोबार [MISSING: company.vat] छ।"),
]
REVIEW = ReviewReport(scores=ReviewScores(completeness=80, relevance=80, compliance=80), verdict="needs_work")


def proposal(language: str = "ne", sections=None) -> Proposal:
    sections = sections or (SECTIONS_NE if language == "ne" else [
        ProposalSection(key="cover_letter", title="Cover letter", body="We submit our bid.", grounded=True, cited_chunk_ids=[1]),
        ProposalSection(key="work_plan", title="Work plan", body="1. Mobilise\n2. Gravel", grounded=True, cited_chunk_ids=[2])])
    return Proposal(id=4, tender_id=1, version=2, language=language, sections=sections, review=REVIEW, status="approved",
                    approved_by="Sita Demo Thapa", approved_at=NOW, citation_rate=0.67)


def package(language: str = "ne", rec=R001, sections=None) -> BidPackage:
    return build_package(record=rec, company=COMPANY, proposal=proposal(language, sections),
                         compliance=ComplianceView(items=ITEMS), documents=DOCS, now=NOW, tender_id=1)


# --- attachment checklist ----------------------------------------------------------------------------------------

def rows_by_type(rows):
    return {r.doc_type.value: r for r in rows}


def test_checklist_marks_present_expiring_missing_with_bs_and_ad_dates():
    rows = rows_by_type(build_checklist(R001, ITEMS, DOCS, NOW, "en"))
    assert [r.doc_type.value for r in build_checklist(R001, ITEMS, DOCS, NOW, "en")] == [
        d.doc_type.value for d in R001.required_docs]  # every listed document; statutory ones are all listed here
    reg = rows["company_registration"]
    assert (reg.status, reg.files, reg.document_ids, reg.expires_on, reg.expires_on_bs) == (
        P, ["01_company_registration.txt"], [1], date(2027, 7, 16), "2084-03-32")
    assert reg.rule_id == "NP-REG-DOCS" and reg.compliance_status.value == "pass"
    lic = rows["business_licence"]  # valid at the deadline, lapses before 2027-02-19
    assert lic.status is E and "2083-09-26 BS / 2027-01-10 AD" in lic.note and "2083-11-07 BS / 2027-02-19 AD" in lic.note
    assert rows["licence_renewal"].status is E  # the licence copy is the renewal evidence (FY 2083/84)
    assert rows["pan_vat"].status is P
    tax = rows["tax_clearance"]  # the demo beat: only the expired FY 2081/82 clearance
    assert tax.status is M and tax.files == [] and tax.fiscal_year_bs == "2082/83" and tax.compliance_status.value == "fail"
    assert "2081/82, expired 2083-03-32 BS / 2026-07-16 AD" in tax.note and "F/Y 2082/83" in tax.note
    assert "Action: Renew tax clearance for F/Y 2082/83" in tax.note
    decl = rows["self_declaration"]
    assert decl.status is M and "section 3 of this package" in decl.note and "Action" not in decl.note
    bid = rows["bid_security"]
    assert bid.status is M and "Rs 4,46,000" in bid.note
    jv = rows["jv_agreement"]
    assert jv.status is M and not jv.mandatory and jv.document.endswith("(optional)")


def test_checklist_present_when_everything_is_valid_and_statutory_rows_are_added():
    fresh = [*DOCS[:2], DocumentMeta(doc_type="tax_return", chunk_id=11, document_id=7, filename="tax_return_2082-83.pdf",
                                     fiscal_year_bs="2082/83"),
             DocumentMeta(doc_type="business_licence", chunk_id=9, document_id=6, filename="lic.txt", expiry=date(2027, 7, 16))]
    rec = record(required_docs=[RequiredDoc(doc_type="tax_clearance", description="Tax clearance F/Y 2082/83",
                                            fiscal_year_bs="2082/83")], bid_validity_days=90)
    rows = build_checklist(rec, [], fresh, NOW, "en")
    assert [(r.doc_type.value, r.status) for r in rows] == [
        ("tax_clearance", P),  # a filed-return proof for that FY satisfies it (PPR Rule 40(2)(d))
        ("company_registration", P), ("business_licence", P), ("pan_vat", P), ("self_declaration", M)]
    assert rows[1].document == "Company registration certificate (statutory, PPR Rule 40(2))"
    assert all(r.rule_id for r in rows) and rows[0].compliance_status is None  # no compliance report given


def test_checklist_unreadable_and_absent_documents():
    scanned = [DocumentMeta(doc_type="company_registration", chunk_id=None, document_id=1, filename="scan.pdf")]
    rows = rows_by_type(build_checklist(record(), [], scanned, NOW, "en"))
    assert rows["company_registration"].status is M and "not readable" in rows["company_registration"].note
    assert rows["pan_vat"].note.startswith("Not on file.")


def test_checklist_nepali_notes_use_devanagari_dates():
    tax = rows_by_type(build_checklist(R001, ITEMS, DOCS, NOW, "ne"))["tax_clearance"]
    assert "२०८३-०३-३२ / २०२६-०७-१६" in tax.note and tax.status is M


# --- package contents --------------------------------------------------------------------------------------------

def test_package_has_cover_letter_bid_form_declarations_sections_and_checklist_in_order():
    pkg = package("en")
    h1 = [b.text for b in pkg.blocks if b.kind == "h1"]
    assert h1 == ["1. Cover letter", "2. Bid form (skeleton)", "3. Declarations", "4. Technical proposal",
                  "5. Attachment checklist"]
    text = pkg.text
    assert "Internal cover sheet: not part of the bid" in text and "Tenderly never submits bids" in text
    assert "We submit our bid." in text  # the proposal's cover_letter section is the letter body
    assert text.count("We submit our bid.") == 1  # ...and is not repeated in the technical proposal
    assert "total bid price of NPR ____________" in text  # never a price
    assert "valid for 90 days from the bid submission deadline (5 Mangsir 2083 BS (21 November 2026 AD))" in text
    assert "We enclose a bid security of Rs 4,46,000" in text and "performance security the bidding document requires" in text
    assert "not been blacklisted or debarred" in text and "PPR Rule 40(2)(e)" in text
    assert "Sources: knowledge-base chunks 2." in text
    assert "Needs work" in text and "Sita Demo Thapa" in text
    [checklist] = [b for b in pkg.blocks if b.kind == "table" and b.rows[0][0] == "#"]
    assert checklist.rows[0] == ["#", "Document", "Status", "File(s)", "Valid until (BS / AD)", "Notes"]
    assert len(checklist.rows) == 1 + len(pkg.checklist)


def test_nepali_package_text_and_flags():
    pkg = package("ne")
    text = pkg.text
    assert "१. आवरण पत्र" in text and "३. स्वघोषणा" in text and "५. संलग्न कागजातको सूची" in text
    assert "५ मंसिर २०८३ (२०२६-११-२१)" in text  # deadline in BS and AD, Devanagari digits
    assert "स्रोत: ज्ञान-आधार खण्ड ७, ९।" in text and "समीक्षा सङ्केत: ungrounded।" in text
    assert "[MISSING: company.vat]" in text  # unresolved markers stay visible, never filled in


def test_markdown_blocks():
    blocks = markdown_blocks("Intro line\ncontinues **bold** and _it_.\n\n- one\n* two\n1. first\n२) second\n## Head")
    assert [b.kind for b in blocks] == ["p", "bullet", "bullet", "number", "number", "p"]
    assert blocks[0].text == "Intro line continues bold and it."
    assert [(i.text, i.bold, i.italic) for i in blocks[0].inlines][1:4] == [("bold", True, False), (" and ", False, False),
                                                                           ("it", False, True)]
    assert markdown_blocks("[UNKNOWN: fact.road_value] and company_profile")[0].inlines[0].italic is False


# --- PDF: Devanagari rendering + extraction -------------------------------------------------------------------------

@pytest.fixture(scope="module")
def nepali_pdf() -> tuple[BidPackage, bytes]:
    pkg = package("ne")
    return pkg, render_pdf(pkg, generated=NOW)


def test_pdf_devanagari_extracts_as_the_same_unicode(nepali_pdf):
    pkg, data = nepali_pdf
    text = norm(pdf_text(data))
    assert not PUA.search(text) and "\x00" not in text and "�" not in text  # no glyph codes or missing glyphs
    words = set(DEVANAGARI_WORD.findall(unicodedata.normalize("NFC", pkg.text)))
    assert len(words) > 150
    assert words <= set(DEVANAGARI_WORD.findall(text)), sorted(words - set(DEVANAGARI_WORD.findall(text)))[:10]
    for phrase in ("क्षेत्रीय सडक निर्माण", "प्राविधिक प्रस्ताव", "श्री प्रमुख प्रशासकीय अधिकृतद्वारा प्रमाणित",
                   "संलग्न कागजातको सूची", "२०८३-०३-३२ / २०२६-०७-१६"):  # conjuncts, reph, ि matra, danda dates
        assert norm(phrase) in text, phrase


def test_pdf_is_shaped_with_embedded_noto_fonts(nepali_pdf):
    _, data = nepali_pdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    fonts = {f[3].split("+")[-1]: f[1] for page in doc for f in page.get_fonts()}
    assert fonts.get("NotoSansDevanagari-Regular") == "ttf" and fonts.get("NotoSansDevanagari-Bold") == "ttf"
    assert fonts.get("NotoSans-Regular") == "ttf"
    # Shaping happened: conjunct/reordered clusters are drawn as non-character glyphs and carry /ActualText.
    content = b"".join(page.read_contents() for page in doc)
    assert content.count(b"/ActualText") > 20
    # Without the ActualText layer the same page would extract glyph codes: proves the spans are what fixes it.
    raw = "".join(page.get_text(flags=pymupdf.TEXTFLAGS_TEXT | pymupdf.TEXT_IGNORE_ACTUALTEXT) for page in doc)
    assert PUA.search(raw)


SEED_FILES = sorted(p for p in (ROOT / "data" / "seed_tenders").glob("*.txt")
                    if DEVANAGARI_WORD.search(p.read_text(encoding="utf-8")))


@pytest.mark.parametrize("path", SEED_FILES, ids=lambda p: p.name[:3])
def test_every_devanagari_seed_notice_round_trips_through_the_pdf_exactly(path):
    notice = path.read_text(encoding="utf-8")
    pkg = BidPackage(language="ne", title=path.name, file_stem="t", checklist=[],
                     blocks=[para(line) for line in notice.splitlines() if line.strip()])
    assert norm(pdf_text(render_pdf(pkg, generated=NOW), body_only=True)) == norm(notice)


@pytest.mark.parametrize("words", [
    "(ऐच्छिक) स्वेच्छिक इच्छिक उच्छिष्ट विच्छेद",  # half-form + ligature with the ि matra: one span per glyph
    "कार्यालय पूर्वाधार निर्माण दर्ता फर्निचर र्‍य",  # reph
    "पालिका लिने कि की कु कृ हिन्दी स्त्री श्री क्षेत्रीय द्ध ह्व ङ्क ट्र सम्झिन्छ अधिकृतद्वारा",
])
def test_hard_devanagari_clusters_extract_exactly(words):
    pkg = BidPackage(language="ne", title="t", file_stem="t", checklist=[], blocks=[para(words)])
    assert norm(pdf_text(render_pdf(pkg, generated=NOW), body_only=True)) == norm(words)


def test_long_devanagari_words_in_narrow_cells_still_extract():
    pkg = BidPackage(language="ne", title="t", file_stem="t", checklist=[], blocks=[
        Block("table", rows=[["क्र.सं.", "कागजात", "x", "y", "z", "प्रमाणपत्रहरूको"]] * 2,
              widths=(0.04, 0.2, 0.2, 0.2, 0.3, 0.06))])
    text = norm(pdf_text(render_pdf(pkg, generated=NOW)))
    assert not PUA.search(text) and "प्रमाणपत्र" in text.replace(" ", "")


def test_english_pdf_has_no_ligature_glyphs():
    text = pdf_text(render_pdf(package("en"), generated=NOW))
    assert "ﬁ" not in text and "fi" in text and "Attachment checklist" in text


# --- DOCX ----------------------------------------------------------------------------------------------------------

def test_docx_keeps_devanagari_and_sets_complex_script_fonts():
    pkg = package("ne")
    doc = Document(io.BytesIO(render_docx(pkg, generated=NOW)))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for phrase in ("क्षेत्रीय सडक निर्माण र कल्भर्ट", "३. स्वघोषणा", "प्राविधिक प्रस्ताव"):
        assert phrase in text
    runs = [r for p in doc.paragraphs for r in p.runs if DEVANAGARI_WORD.search(r.text)]
    assert runs and all(r._element.rPr.find(qn("w:rFonts")).get(qn("w:cs")) == "Noto Sans Devanagari" for r in runs)
    bold = [r for r in runs if r.bold]
    assert bold and all(r._element.rPr.find(qn("w:bCs")) is not None for r in bold)  # Devanagari bold is w:bCs
    for p in doc.paragraphs:  # every rPr child in schema order (Word rejects out-of-order elements)
        for r in p.runs:
            tags = [c.tag.split("}")[1] for c in (r._element.rPr if r._element.rPr is not None else [])]
            assert tags == sorted(tags, key=_RPR_ORDER.index)
    assert doc.core_properties.title.startswith("बोलपत्र पेस गर्ने प्याकेज")


# --- the ZIP route -------------------------------------------------------------------------------------------------

def test_package_route_zips_pdf_docx_checklist_audit_and_attachments(api, db_session, pg_url, make_client,  # noqa: F811
                                                                     session_factory):
    tender = drafted_001(pg_url, make_client, db_session, session_factory, ReviewFake(mention_security=True))
    from agents.review import run_review
    run_review(db_session, db_client(session_factory, ReviewFake(mention_security=True)), tender, settings=settings_(),
               now=NOW)
    [row] = versions(db_session, tender.id)
    assert api.get(f"/proposals/{row.id}/package").status_code == 409
    api.post(f"/proposals/{row.id}/open")
    assert api.post(f"/proposals/{row.id}/approve", json={"approved_by": "Sita Demo Thapa"}).status_code == 200

    resp = api.get(f"/proposals/{row.id}/package")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/zip"
    assert f'filename="tenderly_tender{tender.id}_proposal_v1_package.zip"' in resp.headers["content-disposition"]
    archive = zipfile.ZipFile(io.BytesIO(resp.content))
    names = archive.namelist()
    stem = f"tender{tender.id}_proposal_v1"
    assert {f"01_{stem}.pdf", f"01_{stem}.docx", "02_attachment_checklist.csv", "03_audit_summary.json"} <= set(names)

    # The checklist: the demo's expired FY 2081/82 clearance is missing, registration present (BS + AD dates).
    rows = list(csv.DictReader(io.StringIO(archive.read("02_attachment_checklist.csv").decode("utf-8-sig"))))
    by_type = {r["doc_type"]: r for r in rows}
    assert by_type["tax_clearance"]["status"] == "missing" and "2083-03-32" in by_type["tax_clearance"]["note"].translate(
        str.maketrans("०१२३४५६७८९", "0123456789"))
    assert by_type["company_registration"]["status"] == "present"
    assert by_type["company_registration"]["expires_on_bs"] == "2084-03-32"

    # Attachments: the KB originals of present/expiring rows, byte for byte; never the expired clearance.
    attached = [n for n in names if n.startswith("attachments/")]
    assert any(n.endswith("01_company_registration.txt") for n in attached)
    assert not any("tax_clearance" in n for n in attached)
    reg = db_session.scalar(select(DocumentRow).where(DocumentRow.filename == "01_company_registration.txt"))
    assert archive.read(f"attachments/{reg.id}_01_company_registration.txt") == reg.content

    audit = json.loads(archive.read("03_audit_summary.json"))
    assert audit["submitted"] is False and audit["approval"]["approved_by"] == "Sita Demo Thapa"
    assert audit["review"]["verdict"] == "pass" and audit["proposal"]["version"] == 1
    assert any(r["model"] == "human" and r["stage"] == "review" for r in audit["agent_runs"])
    assert {"compliance", "proposal", "review"} <= {r["stage"] for r in audit["agent_runs"]}
    for f in audit["files"]:  # the listing matches the archive
        assert hashlib.sha256(archive.read(f["name"])).hexdigest() == f["sha256"]
    assert {f["name"] for f in audit["files"]} == set(names) - {"03_audit_summary.json"}

    text = norm(pdf_text(archive.read(f"01_{stem}.pdf")))  # 001 is a mixed notice: a Nepali package
    assert "आवरण पत्र" in text and "संलग्न कागजातको सूची" in text and not PUA.search(text)
    assert Document(io.BytesIO(archive.read(f"01_{stem}.docx"))).paragraphs

    db_session.refresh(row)
    assert build(db_session, row, now=NOW + timedelta(days=1)).content != resp.content  # dated by the clock given
    assert api.get("/proposals/99999/package").status_code == 404
