"""DOCX export of the bid package (python-docx), the editable copy (spec section 10.1).

Word shapes Devanagari itself, so the text is stored as plain Unicode. Every run names Noto Sans for Latin and
Noto Sans Devanagari for complex script (`w:rFonts w:cs`), with the complex-script language set to ne-NP. When the
fonts are not installed, Word falls back to a system Devanagari font (e.g. Nirmala UI). The fonts are not embedded
in the DOCX; the PDF is the fixed-layout copy.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from export.content import BidPackage, Block, Inline

LATIN_FONT, DEVANAGARI_FONT = "Noto Sans", "Noto Sans Devanagari"


# w:rPr children in schema order (ECMA-376 Part 1, 17.3.2.28); Word rejects out-of-order elements.
_RPR_ORDER = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
              "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
              "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
              "lang", "eastAsianLayout", "specVanish", "oMath")


def _put(rpr, tag: str, **attrs: str) -> None:
    """Set (or create, in schema order) one w:rPr child."""
    element = rpr.find(qn(f"w:{tag}"))
    if element is None:
        element = OxmlElement(f"w:{tag}")
        later = {qn(f"w:{t}") for t in _RPR_ORDER[_RPR_ORDER.index(tag) + 1:]}
        successor = next((child for child in rpr if child.tag in later), None)
        if successor is not None:
            successor.addprevious(element)
        else:
            rpr.append(element)
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), value)


def _script(rpr, *, bold: bool = False, italic: bool = False, size: float | None = None) -> None:
    """Latin + complex-script fonts, complex-script bold/italic/size (Devanagari uses the *Cs variants), language."""
    _put(rpr, "rFonts", ascii=LATIN_FONT, hAnsi=LATIN_FONT, cs=DEVANAGARI_FONT)
    if bold:
        _put(rpr, "bCs")
    if italic:
        _put(rpr, "iCs")
    if size:
        _put(rpr, "szCs", val=str(round(size * 2)))
    _put(rpr, "lang", bidi="ne-NP")


def _runs(paragraph, parts: list[Inline], *, size: float | None = None, grey: bool = False) -> None:
    for part in parts:
        run = paragraph.add_run(part.text)
        run.bold = part.bold or None
        run.italic = part.italic or None
        if size:
            run.font.size = Pt(size)
        if grey:
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        _script(run._element.get_or_add_rPr(), bold=part.bold, italic=part.italic, size=size)


def _table(doc, block: Block) -> None:
    header = len(block.rows[0]) > 2
    table = doc.add_table(rows=len(block.rows), cols=len(block.rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(block.rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            _runs(cell.paragraphs[0], [Inline(value, bold=(header and r == 0) or (not header and c == 0))], size=8.5)
    doc.add_paragraph()


def render_docx(package: BidPackage, *, generated: datetime) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    _script(normal.element.get_or_add_rPr(), size=10.5)
    props = doc.core_properties
    props.title, props.subject = package.title, package.file_stem
    props.author, props.comments = "Tenderly", "Prepared with Tenderly for review and signature; not submitted."
    props.created = generated
    props.language = "ne-NP" if package.language == "ne" else "en"

    for block in package.blocks:
        if block.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif block.kind == "table":
            _table(doc, block)
        elif block.kind in ("title", "h1", "h2"):
            heading = doc.add_heading(level={"title": 0, "h1": 1, "h2": 2}[block.kind])
            _runs(heading, block.inlines)
        elif block.kind == "bullet":
            _runs(doc.add_paragraph(style="List Bullet"), block.inlines)
        elif block.kind == "number":
            _runs(doc.add_paragraph(style="List Number"), block.inlines)
        elif block.kind == "note":
            _runs(doc.add_paragraph(), block.inlines, size=8.5, grey=True)
        else:
            _runs(doc.add_paragraph(), block.inlines)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
