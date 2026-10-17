"""Regenerate the binary demo documents in data/demo_company/ (DOCX + PDF).

All content is FICTIONAL. PDFs are English-only: PyMuPDF-written Devanagari
loses conjunct mappings on extraction (see PROGRESS.md), so Devanagari
lives in the TXT and DOCX documents.

    cd api && uv run python -m scripts.build_demo_docs
"""

from pathlib import Path

import docx
import pymupdf

OUT = Path(__file__).resolve().parents[2] / "data" / "demo_company"
BANNER = "FICTIONAL DEMO DOCUMENT - NOT A REAL CERTIFICATE (Tenderly demo data)"
CSS = "body {font-family: sans-serif; font-size: 10.5pt;} h1 {font-size: 16pt;} h2 {font-size: 12.5pt;} td, th {border: 1px solid #999; padding: 3px;} .banner {color: #b00; font-weight: bold;}"


def write_pdf(path: Path, html: str) -> None:
    story = pymupdf.Story(f'<p class="banner">{BANNER}</p>' + html, user_css=CSS)
    writer = pymupdf.DocumentWriter(str(path))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


def write_docx(path: Path, title: str, paragraphs: list[str], table: list[list[str]] | None = None) -> None:
    document = docx.Document()
    document.add_paragraph(BANNER)
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        t.style = "Table Grid"
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.core_properties.author = "Tenderly demo data (fictional)"
    document.save(str(path))


def pan_vat() -> None:
    write_docx(
        OUT / "02_pan_vat_certificate.docx",
        "स्थायी लेखा नम्बर (PAN) तथा मूल्य अभिवृद्धि कर (VAT) दर्ता प्रमाणपत्र",
        [
            "PAN / VAT Registration Certificate",
            "करदाताको नाम: अन्नपूर्ण डेमो निर्माण सेवा प्रा. लि.",
            "Taxpayer name: Annapurna Demo Nirman Sewa Pvt. Ltd.",
            "स्थायी लेखा नम्बर (PAN): ९९९१२३४५६   PAN: 999123456",
            "मू.अ.कर दर्ता नं. (VAT): ९९९१२३४५६   VAT No.: 999123456",
            "दर्ता मिति: २०७१/१२/०५   Date of registration: 2071-12-05 BS",
            "ठेगाना: पोखरा-८, कास्की   Address: Pokhara-8, Kaski",
            "कारोबारको प्रकृति: निर्माण व्यवसाय   Nature of business: Construction",
            "Issuing office: Inland Revenue Office, Pokhara (fictional demo record)",
            "This certificate has no expiry date.",
        ],
    )


def experience() -> None:
    write_docx(
        OUT / "04_experience_certificate.docx",
        "कार्य सम्पन्न प्रमाणपत्र / Work Completion Certificate",
        [
            "Issued by: Demo Rural Municipality, Office of the Municipal Executive, Syangja, Gandaki Province (fictional).",
            "प्रमाणित गरिन्छ कि अन्नपूर्ण डेमो निर्माण सेवा प्रा. लि. ले तल उल्लिखित कार्य सन्तोषजनक रूपमा सम्पन्न गरेको छ।",
            "This is to certify that Annapurna Demo Nirman Sewa Pvt. Ltd. has satisfactorily completed the works below.",
            "Scope: upgrading of the Ward 3 to Ward 5 rural road (4.1 km) to gravel standard, with 3.2 km of "
            "stone-masonry side drains, six hume-pipe culverts (900 mm) and 180 m of gabion retaining wall. "
            "Works were completed within the contract period with no liquidated damages.",
            "The contractor's site engineer and quality-control records were found satisfactory by the municipal technical section.",
        ],
        table=[
            ["Item", "Detail"],
            ["Contract No.", "DRM/NCB/W/07/2080-81"],
            ["Procurement method", "National competitive bidding (NCB)"],
            ["Contract amount (incl. VAT)", "रु. १,८५,४०,००० / Rs. 1,85,40,000"],
            ["Final bill amount", "Rs. 1,83,96,500"],
            ["Start date", "2080-09-15 BS"],
            ["Completion date", "2081-12-20 BS"],
            ["Defects liability period", "1 year, ended 2082-12-20 BS"],
        ],
    )


def tax_clearance() -> None:
    write_pdf(
        OUT / "03_tax_clearance_2081-82.pdf",
        """
<h1>Tax Clearance Certificate</h1>
<p>Fiscal year: <b>2081/82</b></p>
<p>Taxpayer: Annapurna Demo Nirman Sewa Pvt. Ltd.<br/>PAN: 999123456<br/>Address: Pokhara-8, Kaski</p>
<p>This is to certify that the above taxpayer has filed its income tax return for fiscal year 2081/82
and has no outstanding tax liability for that year as of the date of issue.</p>
<p>Date of issue: 2082-04-18 BS</p>
<p><b>Validity:</b> this certificate may be used for public procurement until the end of
fiscal year 2082/83, i.e. until 2083-03-32 BS (Asar end 2083).</p>
<p>Issuing office: Inland Revenue Office, Pokhara (fictional demo record).</p>
""",
    )


def company_profile() -> None:
    projects = [
        ("2081/82", "Ward 3-5 rural road upgrading (4.1 km) with drains and culverts", "Demo Rural Municipality, Syangja", "1,85,40,000"),
        ("2081/82", "Two-storey ward office building (RCC frame, 310 sq m)", "Demo Municipality, Tanahun", "96,75,000"),
        ("2080/81", "Blacktopping (DBST) of 2.2 km urban road", "Pokhara area ward office (demo)", "1,42,10,000"),
        ("2080/81", "Gabion and stone-masonry retaining walls after monsoon landslides", "Demo Rural Municipality, Parbat", "48,60,000"),
        ("2079/80", "School block (4 classrooms) with toilets and water supply", "Demo Secondary School, Lamjung", "63,20,000"),
        ("2079/80", "Irrigation canal lining (1.6 km)", "Demo Rural Municipality, Kaski", "54,90,000"),
    ]
    rows = "".join(f"<tr><td>{fy}</td><td>{name}</td><td>{client}</td><td>{amount}</td></tr>" for fy, name, client, amount in projects)
    write_pdf(
        OUT / "05_company_profile.pdf",
        f"""
<h1>Annapurna Demo Nirman Sewa Pvt. Ltd. - Company Profile</h1>
<p>Pokhara Metropolitan City-8, Kaski, Gandaki Province, Nepal. Registered 2071-11-28 BS (2015).
PAN/VAT 999123456. Construction licence class C (DEMO-C-0421), renewed for FY 2083/84.</p>

<h2>1. Who we are</h2>
<p>Annapurna Demo Nirman Sewa is a construction company that builds rural and urban infrastructure for local
governments in Gandaki Province. Since 2015 we have completed more than 25 contracts for rural municipalities,
municipalities and community schools in Kaski, Tanahun, Syangja, Parbat and Lamjung districts. Our core work is
rural road upgrading, drainage, culverts, slope protection and small public buildings. We bid through national
competitive bidding and sealed quotation, and we partner with Machhapuchchhre Demo Builders in a joint venture for
larger works.</p>

<h2>2. Services</h2>
<ul>
<li>Rural road construction and upgrading: earthwork, sub-base, gravelling, DBST blacktopping.</li>
<li>Drainage: stone-masonry and RCC side drains, hume-pipe and slab culverts, causeways.</li>
<li>Slope protection: gabion walls, stone-masonry retaining walls, bio-engineering after landslides.</li>
<li>Small public buildings: ward offices, school blocks, health posts, with sanitation and water supply.</li>
<li>Irrigation: canal lining and small intake structures.</li>
</ul>

<h2>3. Financial capacity</h2>
<p>Annual turnover (audited): FY 2079/80 Rs. 4,20,00,000; FY 2080/81 Rs. 5,10,00,000; FY 2081/82 Rs. 5,80,00,000.
Average annual turnover over the last three fiscal years is Rs. 5,03,33,333. The company holds a working-capital
credit line of Rs. 1,50,00,000 with a commercial bank and can furnish bid and performance securities as bank guarantees.</p>

<h2>4. Completed projects (selected)</h2>
<table>
<tr><th>FY (BS)</th><th>Project</th><th>Client</th><th>Contract value (NPR)</th></tr>
{rows}
</table>
<p>All six projects were completed on time. Completion certificates are available on request.</p>

<h2>5. Key personnel</h2>
<ul>
<li>Ram Demo Gurung - Managing Director, B.E. Civil, 14 years of road and building construction experience.</li>
<li>Er. Kiran Demo Adhikari - Project Manager, M.Sc. Construction Management, Nepal Engineering Council registered, 11 years.</li>
<li>Er. Anita Demo Pariyar - Site Engineer, B.E. Civil, 6 years, specialises in drainage and slope protection.</li>
<li>Two Sub-Engineers (Diploma in Civil Engineering), 5 and 8 years of site experience.</li>
<li>One Quantity Surveyor and one Accountant; 35 to 60 skilled and unskilled workers depending on season.</li>
</ul>

<h2>6. Equipment owned</h2>
<ul>
<li>2 hydraulic excavators (JCB-class backhoe loaders)</li>
<li>1 vibratory roller, 10 tonne</li>
<li>3 tipper trucks, 10 cubic metre</li>
<li>1 water tanker, 2 concrete mixers, 2 plate compactors</li>
<li>Survey equipment: 1 total station, 2 auto levels</li>
</ul>

<h2>7. Quality, safety and environment</h2>
<p>Every site keeps a daily work diary, material test records (aggregate gradation, cube tests for concrete) and
a measurement book checked by the client's engineer. Workers receive personal protective equipment and a
toolbox safety talk each week. We follow the client's environmental management plan, manage spoil disposal away
from streams, and restore borrow pits after use.</p>

<h2>8. Approach to rural road projects</h2>
<p>We start with a joint site survey with the ward office, confirm alignment and drainage outlets, and schedule
earthwork outside the monsoon. Drains and culverts are built before gravelling so the road base stays dry.
Local labour is hired through the ward wherever possible, and progress is reported to the municipality every
two weeks with photographs and measurement sheets.</p>

<h2>9. Contact</h2>
<p>Pokhara-8, Kaski. Phone: +977-61-000000 (demo). Email: info@annapurna-demo.example</p>
""",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    pan_vat()
    tax_clearance()
    experience()
    company_profile()
    print(f"wrote demo documents to {OUT}")


if __name__ == "__main__":
    main()
